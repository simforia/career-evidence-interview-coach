import io, json
from datetime import datetime
import streamlit as st
from .core import BEHAVIORAL, extract_text, assess_job, targeted_resume, add_log, build_practice_questions, evaluate_practice_answer

try:
    from docx import Document
except Exception:
    Document = None

try:
    from streamlit_mic_recorder import speech_to_text
except Exception:
    speech_to_text = None

def _save(client, user_id, workspace, save_fn):
    save_fn(client, user_id, workspace)
    st.success("Saved.")

def render_app(client, user_id, workspace, save_fn):
    with st.sidebar:
        st.header("Career Workspace")
        if st.button("Sign out", use_container_width=True):
            try:
                client.auth.sign_out()
            except Exception:
                pass
            for key in ["authenticated","user_id","workspace","supabase_client"]:
                st.session_state.pop(key, None)
            st.rerun()

        profiles = workspace.get("interview_profiles", {})
        if profiles:
            names = list(profiles)
            current = workspace.get("selected_profile") if workspace.get("selected_profile") in profiles else names[0]
            selected = st.selectbox("Interview profile", names, index=names.index(current))
            workspace["selected_profile"] = selected

        page = st.radio("Go to", [
            "Dashboard","Master Career Record","Strengths & Gaps",
            "STAR Evidence Builder","Job Review","Practice Sessions",
            "Resume Builder","Workspace Backup"
        ])

    st.title("Career Evidence & Interview Coach")
    st.caption("Evidence-first: every resume claim should be supportable in an interview.")

    if page == "Dashboard":
        a,b,c,d = st.columns(4)
        a.metric("Verified facts", sum(1 for x in workspace.get("master_facts",[]) if x.get("verified")))
        b.metric("STAR stories", sum(1 for x in workspace.get("stories",[]) if x.get("situation") and x.get("result")))
        c.metric("Jobs reviewed", len(workspace.get("jobs",[])))
        d.metric("Practice sessions", len(workspace.get("sessions",[])))
        st.info("Keep environments and experience levels exact. Do not inflate personal or lab work into professional experience.")

        profiles = workspace.get("interview_profiles", {})
        selected = workspace.get("selected_profile")
        if selected in profiles:
            profile = profiles[selected]
            st.subheader(selected)
            st.write(f"**Company:** {profile.get('company','')}  |  **Location:** {profile.get('location','')}")
            st.write(f"**Framework:** {profile.get('framework','')}")
            if profile.get("resume"):
                st.write(f"**Resume:** {profile['resume']}")
            if profile.get("apply"):
                st.link_button("Open application", profile["apply"])
            for item in profile.get("focus",[]):
                st.write(f"• {item}")
        else:
            st.warning("No saved interview profile is loaded yet. Add or review a job to begin building the workspace.")

    elif page == "Master Career Record":
        st.subheader("Master Career Record")
        st.caption("Add only facts the candidate can defend with specifics.")
        with st.expander("Add verified fact"):
            category = st.selectbox("Category", ["Experience","Networking","Cabling","Hardware","Storage","Systems","Security","Education","Certification","Tool","Project","Other"])
            fact = st.text_area("Fact / capability")
            source = st.text_input("Evidence source")
            if st.button("Add fact") and fact.strip():
                workspace.setdefault("master_facts",[]).append({
                    "category":category,"fact":fact.strip(),"source":source or "User-entered","verified":True
                })
                add_log(workspace,"master_fact","Added verified fact",fact.strip())
                _save(client,user_id,workspace,save_fn)

        for i,item in enumerate(workspace.get("master_facts",[])):
            with st.expander(f"{item.get('category','Other')} — {item.get('fact','')[:90]}"):
                st.write(item.get("fact",""))
                st.caption(f"Source: {item.get('source','')}")
                item["verified"] = st.checkbox("Verified / safe for resume use", value=item.get("verified",False), key=f"verified_{i}")

    elif page == "Strengths & Gaps":
        st.subheader("Strengths & Gaps")
        left,right = st.columns(2)
        with left:
            st.markdown("### Strengths")
            for item in workspace.get("strengths",[]):
                if item.get("type") == "strength":
                    st.write(f"**{item.get('area','')}** — {item.get('level','')}\n\n{item.get('evidence','')}")
        with right:
            st.markdown("### Gaps")
            for item in workspace.get("strengths",[]):
                if item.get("type") == "gap":
                    st.write(f"**{item.get('area','')}** — {item.get('level','')}\n\n{item.get('evidence','')}")

        with st.expander("Add item"):
            kind = st.radio("Type", ["strength","gap"], horizontal=True)
            area = st.text_input("Area")
            level = st.text_input("Level")
            evidence = st.text_area("Evidence")
            if st.button("Add strength/gap") and area.strip():
                workspace.setdefault("strengths",[]).append({
                    "area":area.strip(),"level":level.strip(),"evidence":evidence.strip(),"type":kind
                })
                add_log(workspace,"strength_profile",f"Added {kind}",area)
                _save(client,user_id,workspace,save_fn)

    elif page == "STAR Evidence Builder":
        st.subheader("STAR Evidence Builder")
        stories = workspace.setdefault("stories",[])
        if not stories:
            if st.button("Create first story"):
                stories.append({
                    "name":"New story","category":"General","situation":"","task":"","action":"",
                    "result":"","reflection":"","confidence":1,"followups":"","promoted":False
                })
                _save(client,user_id,workspace,save_fn)
                st.rerun()
        else:
            idx = st.selectbox("Story", range(len(stories)), format_func=lambda i: stories[i].get("name","Story"))
            story = stories[idx]
            story["name"] = st.text_input("Story name", story.get("name",""))
            story["category"] = st.text_input("Category", story.get("category",""))
            story["situation"] = st.text_area("Situation", story.get("situation",""))
            story["task"] = st.text_area("Task", story.get("task",""))
            story["action"] = st.text_area("Action — exact steps personally taken", story.get("action",""), height=150)
            story["result"] = st.text_area("Result / verification", story.get("result",""))
            story["reflection"] = st.text_area("Reflection", story.get("reflection",""))
            story["confidence"] = st.slider("Follow-up confidence", 1, 5, int(story.get("confidence",1)))
            story["followups"] = st.text_area("Likely challenges / missing facts", story.get("followups",""))
            if st.button("Save story"):
                stories[idx] = story
                add_log(workspace,"star_story",story["name"],json.dumps(story))
                _save(client,user_id,workspace,save_fn)
            if st.button("Promote to master record", disabled=not all(story.get(k,"").strip() for k in ["situation","action","result"])):
                fact = f"{story['name']}: {story['action']} Result: {story['result']}"
                if not any(x.get("fact")==fact for x in workspace.get("master_facts",[])):
                    workspace.setdefault("master_facts",[]).append({
                        "category":"Project","fact":fact,"source":"STAR Evidence Builder","verified":True
                    })
                _save(client,user_id,workspace,save_fn)

    elif page == "Job Review":
        st.subheader("Upload or paste a job")
        upload = st.file_uploader("Job posting", type=["txt","md","docx","pdf"])
        pasted = st.text_area("Or paste job description", height=220)
        text = pasted.strip() or (extract_text(upload) if upload else "")
        title = st.text_input("Job title")
        company = st.text_input("Company")
        location = st.text_input("Location")
        if st.button("Analyze job") and text:
            score,matched,gaps = assess_job(text,workspace)
            job_id = f"job-{len(workspace.get('jobs',[]))+1}"
            job = {
                "id":job_id,"title":title or "Uploaded Job","company":company or "Unknown",
                "location":location,"text":text,"score":score,"matched":matched,"gaps":gaps,
                "created":datetime.now().isoformat(timespec="seconds")
            }
            workspace.setdefault("jobs",[]).append(job)
            st.session_state.last_job = job_id
            add_log(workspace,"job_review",f"{job['company']} — {job['title']}",f"Evidence signal {score}",job_id=job_id)
            _save(client,user_id,workspace,save_fn)

        if workspace.get("jobs"):
            job = next((j for j in workspace["jobs"] if j["id"]==st.session_state.get("last_job")), workspace["jobs"][-1])
            st.divider()
            st.subheader(f"{job['company']} — {job['title']}")
            a,b,c = st.columns(3)
            a.metric("Evidence-match signal", f"{job['score']}%")
            b.metric("Matched terms", len(job["matched"]))
            c.metric("Potential gaps", len(job["gaps"]))
            st.caption("This is an evidence/keyword signal, not a hiring prediction.")
            st.write("**Supported terms:** " + (", ".join(job["matched"]) or "None detected"))
            st.write("**Potential requirements to verify:** " + (", ".join(job["gaps"]) or "None detected"))
            if st.button("Create practice session"):
                session_id = f"session-{len(workspace.get('sessions',[]))+1}"
                workspace.setdefault("sessions",[]).append({
                    "id":session_id,"job_id":job["id"],"title":f"{job['company']} — {job['title']}",
                    "created":datetime.now().isoformat(timespec="seconds"),"answers":{}
                })
                add_log(workspace,"session","Created practice session",job["title"],job["id"],session_id)
                _save(client,user_id,workspace,save_fn)

    elif page == "Practice Sessions":
        st.subheader("Mock Interview Practice")
        st.caption("Choose a job already in the workspace or a reviewed job, then work through one question at a time.")

        profiles = workspace.get("interview_profiles", {})
        reviewed_jobs = workspace.get("jobs", [])
        targets = []
        for name, profile in profiles.items():
            company = profile.get("company", "")
            label = f"Saved profile — {name}"
            targets.append((label, "profile", name, company))
        for job in reviewed_jobs:
            label = f"Reviewed job — {job.get('company','Unknown')} — {job.get('title','Job')}"
            targets.append((label, "job", job.get("id"), job.get("company","")))

        if targets:
            with st.expander("Start a new mock interview", expanded=not workspace.get("sessions")):
                labels = [x[0] for x in targets]
                selected_label = st.selectbox("Practice for", labels)
                selected_target = next(x for x in targets if x[0] == selected_label)
                mode_new = st.selectbox(
                    "Practice mode",
                    ["Mixed mock","Behavioral / STAR","Resume defense","Job-specific"],
                    help="Mixed mock rotates behavioral, job-specific, and resume-defense questions."
                )
                question_count = st.select_slider("Session length", options=[5,8,10,12], value=8)
                input_new = st.radio(
                    "Answer input",
                    ["Voice","Text"],
                    horizontal=True,
                    help="Voice mode lets you speak the answer and converts it to an editable transcript before scoring."
                )
                if st.button("Start mock interview", type="primary", use_container_width=True):
                    source_type, source_id = selected_target[1], selected_target[2]
                    questions = build_practice_questions(workspace, source_type, source_id, mode_new)[:question_count]
                    session_id = f"session-{len(workspace.get('sessions',[]))+1}"
                    session = {
                        "id": session_id,
                        "source_type": source_type,
                        "source_id": source_id,
                        "job_id": source_id if source_type == "job" else None,
                        "title": selected_label.replace("Saved profile — ","").replace("Reviewed job — ",""),
                        "mode": mode_new,
                        "input_mode": input_new,
                        "questions": questions,
                        "question_index": 0,
                        "answers": {},
                        "evaluations": {},
                        "created": datetime.now().isoformat(timespec="seconds"),
                    }
                    workspace.setdefault("sessions",[]).append(session)
                    add_log(workspace,"session","Started mock interview",f"{session['title']} | {mode_new}",session.get("job_id"),session_id)
                    save_fn(client,user_id,workspace)
                    st.session_state.active_session = session_id
                    st.rerun()
        else:
            st.info("No saved job profiles are available yet. Add a job in Job Review.")

        sessions = workspace.get("sessions",[])
        if sessions:
            st.divider()
            active_id = st.session_state.get("active_session")
            default_idx = next((i for i,s in enumerate(sessions) if s.get("id")==active_id), len(sessions)-1)
            idx = st.selectbox(
                "Practice session",
                range(len(sessions)),
                index=default_idx,
                format_func=lambda i: f"{sessions[i].get('title','Session')} — {sessions[i].get('mode','Practice')}"
            )
            session = sessions[idx]
            st.session_state.active_session = session.get("id")

            source_type = session.get("source_type")
            source_id = session.get("source_id")
            if not source_type:
                source_type = "job" if session.get("job_id") else "profile"
                source_id = session.get("job_id") or workspace.get("selected_profile","")
                session["source_type"] = source_type
                session["source_id"] = source_id

            mode = session.get("mode","Mixed mock")
            questions = session.get("questions") or build_practice_questions(workspace, source_type, source_id, mode)
            session["questions"] = questions
            if not questions:
                st.warning("No practice questions could be generated for this target.")
            else:
                current = min(int(session.get("question_index",0)), len(questions)-1)
                session["question_index"] = current
                completed = len(session.get("evaluations",{}))
                scores = [v.get("score",0) for v in session.get("evaluations",{}).values() if isinstance(v,dict)]
                avg = round(sum(scores)/len(scores)) if scores else 0

                m1,m2,m3 = st.columns(3)
                m1.metric("Progress", f"{completed}/{len(questions)}")
                m2.metric("Current question", f"{current+1}/{len(questions)}")
                m3.metric("Average score", f"{avg}%" if scores else "—")
                st.progress(min(1.0, completed/max(1,len(questions))))

                question = questions[current]
                st.markdown(f"### Question {current+1}")
                st.write(question)
                key = f"q{current}"

                input_options = ["Voice","Text"]
                saved_input = session.get("input_mode","Voice")
                if saved_input not in input_options:
                    saved_input = "Voice"
                input_mode = st.radio(
                    "Answer input",
                    input_options,
                    index=input_options.index(saved_input),
                    horizontal=True,
                    key=f"input_mode_{session.get('id')}_{current}"
                )
                session["input_mode"] = input_mode

                draft_key = f"answer_{session.get('id')}_{current}"
                if draft_key not in st.session_state:
                    st.session_state[draft_key] = session.setdefault("answers",{}).get(key,"")

                if input_mode == "Voice":
                    st.caption("Speak naturally as if the interviewer were in front of you. Microphone permission may be requested by your browser.")
                    if speech_to_text is None:
                        st.warning("Voice transcription is temporarily unavailable. Use Text input for this answer.")
                    else:
                        spoken = speech_to_text(
                            language="en",
                            start_prompt="🎙️ Start answer",
                            stop_prompt="⏹️ Stop & transcribe",
                            just_once=False,
                            use_container_width=True,
                            key=f"voice_{session.get('id')}_{current}"
                        )
                        last_voice_key = f"last_voice_{session.get('id')}_{current}"
                        if spoken and spoken != st.session_state.get(last_voice_key):
                            st.session_state[draft_key] = spoken
                            st.session_state[last_voice_key] = spoken
                    answer = st.text_area(
                        "Transcript — review or correct before submitting",
                        height=240,
                        key=draft_key,
                        placeholder="Your spoken answer will appear here after transcription."
                    )
                    st.caption("Voice transcription is opt-in. The app stores the transcript and scoring data, not the audio recording.")
                else:
                    answer = st.text_area(
                        "Answer",
                        height=240,
                        key=draft_key,
                        placeholder="Answer as if you were speaking to the interviewer. Use a real example and be precise about what you personally did."
                    )

                if st.button("Submit answer for review", type="primary", use_container_width=True):
                    session["answers"][key] = answer
                    evaluation = evaluate_practice_answer(answer, mode, question)
                    session.setdefault("evaluations",{})[key] = evaluation
                    add_log(
                        workspace,
                        "practice_answer",
                        f"Answered question {current+1}",
                        f"Score {evaluation['score']} | {session.get('title','')}",
                        session.get("job_id"),
                        session.get("id")
                    )
                    save_fn(client,user_id,workspace)
                    st.rerun()

                evaluation = session.get("evaluations",{}).get(key)
                if evaluation:
                    st.divider()
                    a,b = st.columns([1,2])
                    a.metric("Answer score", f"{evaluation.get('score',0)}%")
                    a.write(f"**{evaluation.get('label','')}**")
                    with b:
                        st.markdown("**Coach feedback**")
                        for item in evaluation.get("feedback",[]):
                            st.write(f"• {item}")

                    coaching = evaluation.get("coaching", {})
                    if coaching:
                        st.markdown("### Interview-impact coaching")
                        st.write(coaching.get("opening",""))
                        for item in coaching.get("improvements",[]):
                            st.write(f"• {item}")

                        st.markdown("**A stronger way to say the same thing**")
                        st.info(coaching.get("suggested_answer",""))

                        missing = coaching.get("missing",[])
                        if missing:
                            st.warning(
                                "Before using the polished version, add these only if they are true: "
                                + ", ".join(missing) + "."
                            )
                        st.caption(coaching.get("guardrail",""))

                    st.markdown("**Interviewer follow-up**")
                    st.info(evaluation.get("followup","Tell me more about that."))

                    nav1,nav2,nav3 = st.columns(3)
                    if nav1.button("← Previous", disabled=current==0, use_container_width=True):
                        session["question_index"] = max(0,current-1)
                        save_fn(client,user_id,workspace)
                        st.rerun()
                    if nav2.button("Try this answer again", use_container_width=True):
                        session.get("evaluations",{}).pop(key,None)
                        save_fn(client,user_id,workspace)
                        st.rerun()
                    if nav3.button("Next question →", disabled=current>=len(questions)-1, use_container_width=True):
                        session["question_index"] = min(len(questions)-1,current+1)
                        save_fn(client,user_id,workspace)
                        st.rerun()

                weak_counts = {}
                for sess in sessions:
                    for ev in sess.get("evaluations",{}).values():
                        if isinstance(ev,dict):
                            for weak in ev.get("weak_areas",[]):
                                weak_counts[weak] = weak_counts.get(weak,0)+1
                if weak_counts:
                    st.divider()
                    st.markdown("**Recurring practice weaknesses**")
                    names = {
                        "depth":"answer depth",
                        "ownership":"first-person ownership",
                        "technical_detail":"technical detail",
                        "result":"result / verification",
                        "reflection":"reflection / lesson learned",
                        "specificity":"specific evidence",
                    }
                    for weak,count in sorted(weak_counts.items(), key=lambda x:x[1], reverse=True):
                        st.write(f"• {names.get(weak,weak)} — flagged {count} time(s)")

    elif page == "Resume Builder":
        st.subheader("Evidence-only Resume Builder")
        choices = ["No job — general"] + [f"{j['id']} | {j['company']} — {j['title']}" for j in workspace.get("jobs",[])]
        choice = st.selectbox("Target job", choices)
        job = None
        if choice != "No job — general":
            job_id = choice.split(" | ",1)[0]
            job = next(j for j in workspace["jobs"] if j["id"]==job_id)
        text = targeted_resume(job,workspace)
        st.text_area("Generated draft",text,height=500)
        st.download_button("Download TXT",text,"targeted_resume_draft.txt","text/plain",use_container_width=True)
        if Document and st.button("Build DOCX"):
            doc = Document()
            for line in text.splitlines():
                if not line.strip():
                    doc.add_paragraph("")
                elif line.isupper() and not line.startswith("•"):
                    doc.add_heading(line,0 if not doc.paragraphs else 1)
                elif line.startswith("• "):
                    doc.add_paragraph(line[2:],style="List Bullet")
                else:
                    doc.add_paragraph(line)
            bio = io.BytesIO()
            doc.save(bio)
            st.download_button("Download DOCX",bio.getvalue(),"targeted_resume_draft.docx","application/vnd.openxmlformats-officedocument.wordprocessingml.document",use_container_width=True)

    elif page == "Workspace Backup":
        st.subheader("Workspace Backup")
        payload = json.dumps(workspace,indent=2)
        st.download_button("Download complete workspace JSON",payload,"career_workspace.json","application/json",use_container_width=True)
        upload = st.file_uploader("Restore workspace JSON",type=["json"])
        if upload and st.button("Restore"):
            st.session_state.workspace = json.loads(upload.getvalue().decode("utf-8"))
            save_fn(client,user_id,st.session_state.workspace)
            st.rerun()

        st.subheader("Activity log")
        for item in reversed(workspace.get("practice_log",[])[-50:]):
            st.write(f"**{item.get('time')} — {item.get('kind')} — {item.get('title')}**")
            if item.get("body"):
                st.caption(item["body"][:500])

    save_fn(client,user_id,workspace)
